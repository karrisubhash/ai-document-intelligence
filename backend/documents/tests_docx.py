"""DOCX ingestion: detection, expansion guards, extraction and pagination.

The fixtures are built with python-docx rather than checked in as binaries,
so what each test exercises is visible in the test itself.
"""

import io
import zipfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from rest_framework import status

from documents.models import Document, PageContent, ProcessingJob
from documents.processing.context import FileType
from documents.processing.errors import ErrorCode, StageError
from documents.processing.extractors import docx as docx_extractor
from documents.processing.stages.ingestion import detect_file_type
from documents.tests import MEDIA_ROOT, AuthenticatedAPITestCase


def make_docx_bytes(paragraphs=(), tables=(), page_breaks=()):
    """A real .docx in memory.

    `page_breaks` holds indices into `paragraphs` after which an explicit
    page break is inserted.
    """
    import docx

    document = docx.Document()
    for index, text in enumerate(paragraphs):
        paragraph = document.add_paragraph(text)
        if index in page_breaks:
            document.add_page_break()

    for rows in tables:
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for row_index, row in enumerate(rows):
            for cell_index, value in enumerate(row):
                table.cell(row_index, cell_index).text = value

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def write(tmp_path, name, payload):
    path = tmp_path / name
    path.write_bytes(payload)
    return str(path)


class DocxDetectionTests(TestCase):
    def setUp(self):
        import tempfile
        from pathlib import Path

        self.tmp = Path(tempfile.mkdtemp())

    def test_detects_a_real_docx_despite_the_generic_zip_signature(self):
        path = write(self.tmp, "a.docx", make_docx_bytes(["hello"]))
        self.assertEqual(detect_file_type(path), FileType.DOCX)

    def test_a_plain_zip_is_not_mistaken_for_a_docx(self):
        """Every ZIP-based office format shares DOCX's magic bytes."""
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("hello.txt", "not a word document")
        path = write(self.tmp, "archive.zip", buffer.getvalue())

        self.assertIsNone(detect_file_type(path))

    def test_a_zip_named_docx_is_still_not_a_docx(self):
        """The extension is client-supplied and must not decide."""
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("xl/workbook.xml", "<workbook/>")
        path = write(self.tmp, "spreadsheet.docx", buffer.getvalue())

        self.assertIsNone(detect_file_type(path))

    def test_is_docx_returns_false_for_a_corrupt_archive(self):
        path = write(self.tmp, "broken.docx", b"PK\x03\x04 truncated garbage")
        self.assertFalse(docx_extractor.is_docx(path))


class DocxExpansionGuardTests(TestCase):
    """A DOCX is a ZIP, and ZIPs from untrusted users can be bombs."""

    def setUp(self):
        import tempfile
        from pathlib import Path

        self.tmp = Path(tempfile.mkdtemp())

    def test_a_normal_document_passes(self):
        path = write(self.tmp, "ok.docx", make_docx_bytes(["short"]))
        report = docx_extractor.check_expansion(path)

        self.assertGreater(report["uncompressed_bytes"], 0)
        self.assertIn("compression_ratio", report)

    def test_an_absurd_compression_ratio_is_refused_unparsed(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", "A" * 5_000_000)
        path = write(self.tmp, "bomb.docx", buffer.getvalue())

        with self.assertRaises(StageError) as caught:
            docx_extractor.check_expansion(path)

        self.assertEqual(caught.exception.code, ErrorCode.CORRUPT_DOCUMENT)
        self.assertIn("bomb", str(caught.exception).lower())

    def test_the_guard_runs_before_parsing(self):
        """A bomb must be refused without the parser ever seeing it."""
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", "A" * 5_000_000)
        path = write(self.tmp, "bomb2.docx", buffer.getvalue())

        with self.assertRaises(StageError):
            docx_extractor.extract_pages(path)

    def test_corrupt_archive_raises_a_classified_error(self):
        path = write(self.tmp, "corrupt.docx", b"PK\x03\x04 not really a zip")
        with self.assertRaises(StageError) as caught:
            docx_extractor.check_expansion(path)
        self.assertEqual(caught.exception.code, ErrorCode.CORRUPT_DOCUMENT)


class DocxExtractionTests(TestCase):
    def setUp(self):
        import tempfile
        from pathlib import Path

        self.tmp = Path(tempfile.mkdtemp())

    def extract(self, name, **kwargs):
        path = write(self.tmp, name, make_docx_bytes(**kwargs))
        return docx_extractor.extract_pages(path)

    def test_reads_paragraph_text(self):
        pages, _ = self.extract("p.docx", paragraphs=["First line", "Second line"])
        self.assertEqual(len(pages), 1)
        self.assertIn("First line", pages[0])
        self.assertIn("Second line", pages[0])

    def test_reads_table_text_which_paragraphs_alone_would_miss(self):
        pages, metadata = self.extract(
            "t.docx",
            paragraphs=["Fee structure"],
            tables=[[["Course", "Fee"], ["CS101", "5000"]]],
        )
        self.assertEqual(metadata["tables"], 1)
        self.assertIn("CS101", pages[0])
        self.assertIn("5000", pages[0])

    def test_keeps_paragraphs_and_tables_in_document_order(self):
        pages, _ = self.extract(
            "order.docx",
            paragraphs=["Before the table"],
            tables=[[["Inside", "Table"]]],
        )
        body = pages[0]
        self.assertLess(body.index("Before the table"), body.index("Inside"))

    def test_explicit_page_breaks_split_pages(self):
        pages, metadata = self.extract(
            "b.docx",
            paragraphs=["Page one body", "Page two body"],
            page_breaks=(0,),
        )
        self.assertEqual(len(pages), 2)
        self.assertIn("Page one body", pages[0])
        self.assertIn("Page two body", pages[1])
        self.assertEqual(metadata["pagination"], "explicit_breaks")

    def test_pagination_is_labelled_none_when_the_author_declared_no_breaks(self):
        """A DOCX has no rendered page count; saying so is the honest
        alternative to inventing one."""
        _, metadata = self.extract("n.docx", paragraphs=["Only content"])
        self.assertEqual(metadata["pagination"], "none")

    def test_reports_the_extractor_and_its_version(self):
        _, metadata = self.extract("v.docx", paragraphs=["x"])
        self.assertEqual(metadata["extractor"], "python-docx")
        self.assertTrue(metadata["extractor_version"])

    def test_page_count_matches_extracted_pages(self):
        path = write(
            self.tmp,
            "c.docx",
            make_docx_bytes(paragraphs=["a", "b", "c"], page_breaks=(0, 1)),
        )
        self.assertEqual(docx_extractor.page_count(path), 3)


@override_settings(
    MEDIA_ROOT=MEDIA_ROOT,
    CELERY_TASK_ALWAYS_EAGER=True,
    CELERY_TASK_EAGER_PROPAGATES=True,
)
class DocxPipelineTests(AuthenticatedAPITestCase):
    """End to end through the real pipeline, via the API."""

    DOCX_MIME = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    def _post(self, name, payload):
        return self.client.post(
            reverse("document-list"),
            {
                "title": "DOCX upload",
                "file": SimpleUploadedFile(name, payload, content_type=self.DOCX_MIME),
            },
            format="multipart",
        )

    def upload(self, name, payload):
        with self.captureOnCommitCallbacks(execute=True):
            return self._post(name, payload)

    def upload_expecting_failure(self, name, payload):
        """Permanent failures propagate as StageError - the task does not
        retry a document that can never succeed."""
        with self.assertRaises(StageError):
            with self.captureOnCommitCallbacks(execute=True):
                response = self._post(name, payload)
                self.document_id = response.data["id"]
        return self.document_id

    def test_a_docx_is_processed_to_completion(self):
        response = self.upload(
            "regulations.docx",
            make_docx_bytes(paragraphs=["Attendance minimum is 75 percent."]),
        )
        self.assertEqual(response.status_code, status.HTTP_202_ACCEPTED)

        document = Document.objects.get(id=response.data["id"])
        self.assertEqual(document.status, Document.Status.COMPLETED)

    def test_pages_are_recorded_as_docx_text_not_plain_text(self):
        response = self.upload("r.docx", make_docx_bytes(paragraphs=["Body text"]))
        document = Document.objects.get(id=response.data["id"])
        page = PageContent.objects.get(
            content__document_version=document.current_version
        )

        self.assertEqual(page.extraction_method, PageContent.ExtractionMethod.DOCX_TEXT)
        self.assertIn("Body text", page.text)

    def test_docx_never_reaches_ocr(self):
        response = self.upload("r.docx", make_docx_bytes(paragraphs=["Has a text layer"]))
        document = Document.objects.get(id=response.data["id"])
        job = ProcessingJob.objects.get(document_version=document.current_version)
        ocr_event = job.events.get(stage=ProcessingJob.Stage.OCR)

        self.assertEqual(ocr_event.metadata.get("pages_ocred"), 0)

    def test_a_docx_with_no_text_fails_honestly(self):
        """An image-only DOCX yields nothing, and this pipeline does not
        claim to OCR embedded images."""
        document_id = self.upload_expecting_failure(
            "empty.docx", make_docx_bytes(paragraphs=[])
        )
        document = Document.objects.get(id=document_id)
        job = ProcessingJob.objects.get(document_version=document.current_version)

        self.assertEqual(document.status, Document.Status.FAILED)
        self.assertEqual(job.error_code, ErrorCode.EMPTY_EXTRACTION)

    def test_a_zip_pretending_to_be_a_docx_is_rejected(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("xl/workbook.xml", "<workbook/>")

        document_id = self.upload_expecting_failure("fake.docx", buffer.getvalue())
        document = Document.objects.get(id=document_id)
        job = ProcessingJob.objects.get(document_version=document.current_version)

        self.assertEqual(document.status, Document.Status.FAILED)
        self.assertEqual(job.error_code, ErrorCode.UNSUPPORTED_FILE_TYPE)

    def test_each_page_gets_its_own_metadata_dict(self):
        """Sharing one dict across pages would alias them, and document-level
        counts stated per page would read as per-page facts."""
        response = self.upload(
            "multi.docx",
            make_docx_bytes(paragraphs=["one", "two"], page_breaks=(0,)),
        )
        document = Document.objects.get(id=response.data["id"])
        pages = list(
            PageContent.objects.filter(
                content__document_version=document.current_version
            ).order_by("page_number")
        )

        self.assertEqual(len(pages), 2)
        for page in pages:
            self.assertEqual(page.metadata["extractor"], "python-docx")
            # Document-level counts must NOT be restated on each page.
            self.assertNotIn("paragraphs", page.metadata)
            self.assertNotIn("tables", page.metadata)
            self.assertNotIn("uncompressed_bytes", page.metadata)

    def test_document_level_stats_land_on_the_content_row(self):
        response = self.upload(
            "stats.docx",
            make_docx_bytes(
                paragraphs=["body"], tables=[[["a", "b"]]]
            ),
        )
        document = Document.objects.get(id=response.data["id"])
        content = document.current_version.contents.first()

        self.assertEqual(content.metadata["tables"], 1)
        self.assertIn("paragraphs", content.metadata)
        self.assertIn("compression_ratio", content.metadata)

    def test_tables_survive_into_the_stored_content(self):
        response = self.upload(
            "fees.docx",
            make_docx_bytes(
                paragraphs=["Fee schedule"],
                tables=[[["Course", "Fee"], ["CS101", "5000"]]],
            ),
        )
        document = Document.objects.get(id=response.data["id"])
        page = PageContent.objects.get(
            content__document_version=document.current_version
        )
        self.assertIn("CS101", page.text)

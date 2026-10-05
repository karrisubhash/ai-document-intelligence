"""Phase 7: malware scanning, quarantine and rate limiting.

The scanner tests build their fixtures byte by byte rather than checking in
sample files, so what each rule fires on is visible in the test itself - and
so the repository never contains anything that looks like malware.
"""

import io
import zipfile
from unittest.mock import patch

from django.core.cache import cache
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APIClient
from rest_framework.throttling import SimpleRateThrottle

from documents.models import Document, ProcessingJob
from documents.processing.errors import ErrorCode, TransientError
from documents.processing.pipeline import run_pipeline
from documents.processing.scanning import (
    DocumentInfected,
    ScannerNotConfigured,
    ScannerUnavailable,
    Verdict,
)
from documents.processing.scanning import service as scanning_service
from documents.processing.scanning.heuristic import EICAR, HeuristicScanner
from documents.tasks import is_retryable
from documents.tests import MEDIA_ROOT, make_user

SCANNER = "documents.processing.scanning.heuristic.HeuristicScanner"


def write(tmp, name, payload):
    path = tmp / name
    path.write_bytes(payload)
    return str(path)


def pdf_bytes(extra=b""):
    return b"%PDF-1.4\n" + extra + b"\ntrailer\n%%EOF\n"


def ooxml_bytes(parts):
    """A structurally minimal OOXML archive.

    Enough for the scanner, which only reads the archive's name list, but
    NOT a parseable Word document - see real_docx_with for that.
    """
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "<document/>")
        for name, content in parts.items():
            archive.writestr(name, content)
    return buffer.getvalue()


def real_docx_with(parts):
    """A genuinely parseable .docx with extra parts spliced in.

    The scanner only reads the name list, but the pipeline goes on to hand
    the same bytes to python-docx - so a fixture that is merely zip-shaped
    tests the scanner and then breaks extraction.
    """
    import docx

    document = docx.Document()
    document.add_paragraph("Attendance minimum is 75 percent.")
    original = io.BytesIO()
    document.save(original)

    source = zipfile.ZipFile(io.BytesIO(original.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for item in source.infolist():
            archive.writestr(item, source.read(item.filename))
        for name, content in parts.items():
            archive.writestr(name, content)
    return out.getvalue()


class UnknownVerdictScanner:
    name = "weird"
    version = "0.0.1"
    ruleset_version = "x"

    def scan(self, file_path):
        from documents.processing.scanning.base import ScanResult

        return ScanResult(verdict="PROBABLY_FINE", scanner_name=self.name)


class DownScanner:
    """Stands in for a scanner host that is unreachable."""

    name = "down"
    version = "0"
    ruleset_version = "0"

    def scan(self, file_path):
        raise ScannerUnavailable("scanner host unreachable")


# --- The scanner ------------------------------------------------------------


class HeuristicScannerTests(TestCase):
    def setUp(self):
        import tempfile
        from pathlib import Path

        self.tmp = Path(tempfile.mkdtemp())
        self.scanner = HeuristicScanner()

    def scan(self, name, payload):
        return self.scanner.scan(write(self.tmp, name, payload))

    def test_an_ordinary_document_is_clean(self):
        result = self.scan("ok.txt", b"Attendance minimum is 75 percent.\n")
        self.assertEqual(result.verdict, Verdict.CLEAN)
        self.assertEqual(result.findings, ())

    def test_eicar_test_signature_is_detected(self):
        """Proves the scan path runs without anyone needing real malware."""
        result = self.scan("eicar.txt", EICAR)

        self.assertEqual(result.verdict, Verdict.INFECTED)
        self.assertEqual(result.findings[0].rule, "eicar")

    def test_eicar_is_found_even_when_buried_in_a_large_file(self):
        payload = (b"A" * 3_000_000) + EICAR + (b"B" * 1000)
        self.assertEqual(self.scan("big.txt", payload).verdict, Verdict.INFECTED)

    def test_eicar_straddling_a_chunk_boundary_is_still_found(self):
        """The read overlap exists for exactly this case."""
        from documents.processing.scanning.heuristic import CHUNK

        payload = (b"A" * (CHUNK - 20)) + EICAR + b"tail"
        self.assertEqual(self.scan("edge.txt", payload).verdict, Verdict.INFECTED)

    def test_an_executable_disguised_as_a_pdf_is_infected(self):
        result = self.scan("invoice.pdf", b"MZ\x90\x00" + b"\x00" * 200)

        self.assertEqual(result.verdict, Verdict.INFECTED)
        self.assertEqual(result.findings[0].rule, "executable_header")

    def test_an_elf_binary_is_infected(self):
        result = self.scan("report.pdf", b"\x7fELF\x02\x01\x01" + b"\x00" * 100)
        self.assertEqual(result.verdict, Verdict.INFECTED)

    def test_a_shell_script_is_infected(self):
        result = self.scan("notes.txt", b"#!/bin/sh\nrm -rf /\n")
        self.assertEqual(result.verdict, Verdict.INFECTED)

    def test_pdf_javascript_is_suspicious_not_infected(self):
        """Active content is a capability, not proof of intent."""
        result = self.scan("form.pdf", pdf_bytes(b"/JavaScript (app.alert(1))"))

        self.assertEqual(result.verdict, Verdict.SUSPICIOUS)
        self.assertTrue(any(f.rule == "pdf_active_content" for f in result.findings))

    def test_pdf_launch_action_is_flagged(self):
        result = self.scan("x.pdf", pdf_bytes(b"/Launch /F (calc.exe)"))
        self.assertEqual(result.verdict, Verdict.SUSPICIOUS)

    def test_a_plain_pdf_is_clean(self):
        self.assertEqual(self.scan("plain.pdf", pdf_bytes()).verdict, Verdict.CLEAN)

    def test_an_office_macro_is_flagged(self):
        result = self.scan("macro.docx", ooxml_bytes({"word/vbaProject.bin": b"\x00m"}))

        self.assertEqual(result.verdict, Verdict.SUSPICIOUS)
        self.assertTrue(any(f.rule == "ooxml_macro" for f in result.findings))

    def test_a_macro_free_docx_is_clean(self):
        self.assertEqual(
            self.scan("plain.docx", ooxml_bytes({})).verdict, Verdict.CLEAN
        )

    def test_result_records_what_it_scanned_with(self):
        """A clean verdict is only meaningful alongside the rules that
        produced it."""
        result = self.scan("ok.txt", b"hello world")

        self.assertEqual(result.scanner_name, "heuristic-scanner")
        self.assertTrue(result.ruleset_version)
        self.assertFalse(result.metadata["is_antivirus"])

    def test_an_unreadable_file_raises_rather_than_returning_clean(self):
        with self.assertRaises(ScannerUnavailable):
            self.scanner.scan("/nonexistent/path/file.pdf")


# --- The service boundary ---------------------------------------------------


class ScanServiceTests(TestCase):
    def setUp(self):
        scanning_service.reset_scanner_cache()
        self.addCleanup(scanning_service.reset_scanner_cache)

    def test_unconfigured_scanner_raises_not_configured(self):
        with override_settings(DOCUMENT_SCANNER=""):
            with self.assertRaises(ScannerNotConfigured):
                # The path is never resolved. The assertion is that an
                # unconfigured scanner refuses BEFORE touching the
                # filesystem, so a path that does not exist is the point.
                scanning_service.scan_file("/tmp/whatever")  # nosec B108

    def test_an_unknown_verdict_is_never_read_as_clean(self):
        import tempfile
        from pathlib import Path

        path = Path(tempfile.mkdtemp()) / "f.txt"
        path.write_bytes(b"hello")

        with override_settings(DOCUMENT_SCANNER=f"{__name__}.UnknownVerdictScanner"):
            with self.assertRaises(ScannerUnavailable):
                scanning_service.scan_file(str(path))

    def test_cache_is_keyed_on_the_ruleset(self):
        with override_settings(DOCUMENT_SCANNER=SCANNER, DOCUMENT_SCANNER_RULESET="a"):
            first = scanning_service.get_scanner()
        with override_settings(DOCUMENT_SCANNER=SCANNER, DOCUMENT_SCANNER_RULESET="b"):
            second = scanning_service.get_scanner()

        self.assertIsNot(first, second)


# --- The pipeline stage and quarantine --------------------------------------


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class ScanStageTests(TestCase):
    def setUp(self):
        scanning_service.reset_scanner_cache()
        self.addCleanup(scanning_service.reset_scanner_cache)
        self.user = make_user("scan_user", department="Registry")

    def upload(self, payload, name="doc.txt"):
        document = Document.objects.create(
            title="Scan probe", owner=self.user, department="Registry"
        )
        version = document.versions.create(
            version_number=1,
            file=ContentFile(payload, name=name),
            file_name=name,
            file_size=len(payload),
            uploaded_by=self.user,
        )
        document.current_version = version
        document.save(update_fields=["current_version"])
        return document, ProcessingJob.objects.create(document_version=version)

    def test_a_clean_document_passes_the_stage(self):
        document, job = self.upload(b"Attendance minimum is 75 percent.\n")

        with override_settings(DOCUMENT_SCANNER=SCANNER):
            run_pipeline(job.id)

        document.refresh_from_db()
        self.assertEqual(document.status, Document.Status.COMPLETED)
        event = job.events.get(stage=ProcessingJob.Stage.SCANNING)
        self.assertEqual(event.metadata["verdict"], Verdict.CLEAN)

    def test_an_infected_document_is_quarantined_not_failed(self):
        """Spec 10.1 keeps these apart: FAILED invites a retry, QUARANTINED
        is a decision not to touch the file."""
        document, job = self.upload(EICAR, name="eicar.txt")

        with override_settings(DOCUMENT_SCANNER=SCANNER):
            with self.assertRaises(DocumentInfected):
                run_pipeline(job.id)

        document.refresh_from_db()
        job.refresh_from_db()
        self.assertEqual(document.status, Document.Status.QUARANTINED)
        self.assertEqual(job.error_code, ErrorCode.MALWARE_DETECTED)

    def test_a_quarantined_file_cannot_be_downloaded(self):
        """Quarantine stopped the parsing but not the download: every
        department member could pull the flagged bytes as an attachment."""
        document, job = self.upload(EICAR, name="eicar.txt")
        with override_settings(DOCUMENT_SCANNER=SCANNER):
            with self.assertRaises(DocumentInfected):
                run_pipeline(job.id)

        colleague = make_user("scan_colleague", department="Registry")
        client = APIClient()
        client.force_authenticate(colleague)
        url = reverse("document-download", args=[document.id])

        self.assertEqual(client.get(url).status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(
            client.get(url, {"version": 1}).status_code, status.HTTP_403_FORBIDDEN
        )

        from accounts.models import UserProfile

        reviewer = make_user("scan_reviewer", UserProfile.Role.SYSTEM_ADMIN, "")
        client.force_authenticate(reviewer)
        self.assertEqual(client.get(url).status_code, status.HTTP_200_OK)

    def quarantine(self):
        document, job = self.upload(EICAR, name="eicar.txt")
        with override_settings(DOCUMENT_SCANNER=SCANNER):
            with self.assertRaises(DocumentInfected):
                run_pipeline(job.id)
        return document, job

    def test_the_owner_cannot_reprocess_a_quarantined_file(self):
        """The download gate read the NEWEST job, and a new job's error code
        is empty: the owner clicked Reprocess and, until a worker rescanned
        it, every department member could download the flagged file."""
        document, _ = self.quarantine()
        owner = APIClient()
        owner.force_authenticate(self.user)

        response = owner.post(reverse("document-reprocess", args=[document.id]))

        self.assertEqual(response.status_code, status.HTTP_409_CONFLICT)
        self.assertEqual(
            ProcessingJob.objects.filter(
                document_version=document.current_version
            ).count(),
            1,
        )
        document.refresh_from_db()
        self.assertEqual(document.status, Document.Status.QUARANTINED)

    def test_an_administrator_cannot_retry_a_quarantined_job(self):
        from accounts.models import UserProfile

        document, job = self.quarantine()
        admin = make_user(
            "scan_dept_admin", UserProfile.Role.DEPARTMENT_ADMIN, "Registry"
        )
        client = APIClient()
        client.force_authenticate(admin)

        response = client.post(reverse("admin-job-retry", args=[job.id]))

        self.assertEqual(response.status_code, status.HTTP_409_CONFLICT)
        self.assertEqual(
            ProcessingJob.objects.filter(
                document_version=document.current_version
            ).count(),
            1,
        )

    def test_only_a_completed_rerun_lifts_the_quarantine(self):
        """Whatever creates a new job, a queued one does not clear the
        verdict; only a run that completed - and so passed the scan - does."""
        document, _ = self.quarantine()
        version = document.current_version
        colleague = make_user("scan_colleague_2", department="Registry")
        client = APIClient()
        client.force_authenticate(colleague)
        url = reverse("document-download", args=[document.id])

        ProcessingJob.objects.create(document_version=version)
        self.assertEqual(client.get(url).status_code, status.HTTP_403_FORBIDDEN)

        ProcessingJob.objects.create(
            document_version=version, status=ProcessingJob.Status.COMPLETED
        )
        self.assertEqual(client.get(url).status_code, status.HTTP_200_OK)

    def test_an_infected_document_is_never_parsed(self):
        """The whole point of scanning before extraction."""
        document, job = self.upload(EICAR, name="eicar.txt")

        with override_settings(DOCUMENT_SCANNER=SCANNER):
            with self.assertRaises(DocumentInfected):
                run_pipeline(job.id)

        self.assertFalse(
            job.events.filter(stage=ProcessingJob.Stage.TEXT_EXTRACTION).exists()
        )
        self.assertFalse(document.current_version.contents.exists())

    def test_an_executable_disguised_as_a_document_is_quarantined(self):
        """Found live: validation ran first and rejected it as merely an
        unsupported type, throwing away the one fact worth escalating -
        that somebody uploaded a binary."""
        document, job = self.upload(b"MZ" + bytes(304), name="invoice.pdf")

        with override_settings(DOCUMENT_SCANNER=SCANNER):
            with self.assertRaises(DocumentInfected):
                run_pipeline(job.id)

        document.refresh_from_db()
        job.refresh_from_db()
        self.assertEqual(document.status, Document.Status.QUARANTINED)
        self.assertEqual(job.error_code, ErrorCode.MALWARE_DETECTED)

    def test_scanning_runs_before_validation(self):
        """A hostile file must be identified as hostile whether or not it
        also happens to be a type the platform supports."""
        document, job = self.upload(EICAR, name="eicar.xyz")

        with override_settings(DOCUMENT_SCANNER=SCANNER):
            with self.assertRaises(DocumentInfected):
                run_pipeline(job.id)

        stages = list(job.events.values_list("stage", flat=True))
        self.assertIn(ProcessingJob.Stage.SCANNING, stages)
        self.assertNotIn(ProcessingJob.Stage.VALIDATION, stages)

    def test_infection_is_not_retried(self):
        """A file that is malware does not become clean on attempt two."""
        self.assertFalse(is_retryable(DocumentInfected("found something")))

    def test_the_stage_skips_honestly_when_no_scanner_is_configured(self):
        document, job = self.upload(b"ordinary text\n")

        with override_settings(DOCUMENT_SCANNER=""):
            run_pipeline(job.id)

        event = job.events.get(stage=ProcessingJob.Stage.SCANNING)
        self.assertEqual(event.status, "SKIPPED")
        self.assertIn("not scanned", event.message)

    def test_a_suspicious_document_still_processes_but_is_recorded(self):
        payload = real_docx_with({"word/vbaProject.bin": b"\x00macro"})
        document, job = self.upload(payload, name="macro.docx")

        with override_settings(DOCUMENT_SCANNER=SCANNER):
            run_pipeline(job.id)

        document.refresh_from_db()
        event = job.events.get(stage=ProcessingJob.Stage.SCANNING)
        self.assertEqual(document.status, Document.Status.COMPLETED)
        self.assertEqual(event.metadata["verdict"], Verdict.SUSPICIOUS)
        self.assertTrue(event.metadata["findings"])

    def test_scanner_unavailable_is_transient_and_blocks_the_document(self):
        """An unscanned file must not slip through because the scanner is
        down - that is exactly when the control matters."""
        document, job = self.upload(b"text\n")

        with override_settings(DOCUMENT_SCANNER=f"{__name__}.DownScanner"):
            with self.assertRaises(TransientError):
                run_pipeline(job.id)

        document.refresh_from_db()
        self.assertNotEqual(document.status, Document.Status.COMPLETED)


# --- Rate limiting ----------------------------------------------------------


def rates(**overrides):
    """Patch DRF's throttle rates for one test.

    override_settings cannot reach these: SimpleRateThrottle binds
    THROTTLE_RATES = api_settings.DEFAULT_THROTTLE_RATES as a CLASS
    attribute at import time, so changing the setting afterwards leaves the
    already-bound dict in place. Patching that dict is what actually
    exercises the throttle.
    """
    generous = {
        "user": "1000/hour",
        "anon": "1000/hour",
        "login": "1000/minute",
        "login_account": "1000/hour",
        "upload": "1000/hour",
        "search": "1000/hour",
        "chat": "1000/hour",
    }
    generous.update(overrides)
    return patch.dict(SimpleRateThrottle.THROTTLE_RATES, generous)


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class ThrottlingTests(TestCase):
    def setUp(self):
        cache.clear()
        self.addCleanup(cache.clear)
        self.user = make_user("throttle_user", department="Registry")
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def search(self, client=None):
        return (client or self.client).post(
            reverse("search"), {"query": "attendance"}, format="json"
        )

    def test_search_is_throttled_at_its_own_rate(self):
        with rates(search="2/hour"):
            self.assertEqual(self.search().status_code, status.HTTP_200_OK)
            self.assertEqual(self.search().status_code, status.HTTP_200_OK)
            self.assertEqual(
                self.search().status_code, status.HTTP_429_TOO_MANY_REQUESTS
            )

    def test_throttling_one_user_does_not_throttle_another(self):
        """The account is the unit of abuse, not the address."""
        with rates(search="2/hour"):
            for _ in range(3):
                self.search()

            other = make_user("throttle_other", department="Registry")
            client = APIClient()
            client.force_authenticate(other)
            self.assertEqual(self.search(client).status_code, status.HTTP_200_OK)

    def test_chat_has_a_separate_budget_from_search(self):
        """Exhausting the cheap endpoint must not disable the expensive one:
        they are priced separately for a reason."""
        with rates(search="1/hour", chat="5/hour"):
            self.search()
            self.assertEqual(
                self.search().status_code, status.HTTP_429_TOO_MANY_REQUESTS
            )
            response = self.client.post(
                reverse("chat"), {"question": "attendance"}, format="json"
            )
            self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_login_is_throttled_against_credential_guessing(self):
        client = APIClient()
        url = reverse("api-token-auth")

        with rates(login="2/minute"):
            for _ in range(2):
                client.post(
                    url, {"username": "victim", "password": "wrong"}, format="json"
                )
            response = client.post(
                url, {"username": "victim", "password": "wrong"}, format="json"
            )

        self.assertEqual(response.status_code, status.HTTP_429_TOO_MANY_REQUESTS)

    def test_guessing_one_account_does_not_lock_out_another(self):
        """Keying on username alone would let anyone lock a known account
        out; keying on IP alone would let one NAT block a whole office."""
        client = APIClient()
        url = reverse("api-token-auth")

        with rates(login="2/minute"):
            for _ in range(3):
                client.post(
                    url, {"username": "victim", "password": "wrong"}, format="json"
                )
            response = client.post(
                url, {"username": "someone_else", "password": "wrong"}, format="json"
            )

        self.assertNotEqual(response.status_code, status.HTTP_429_TOO_MANY_REQUESTS)

    def test_whitespace_around_the_username_does_not_reset_the_limit(self):
        """The serializer strips the username before authenticating, so
        "victim " IS victim. Keyed on the raw value, every padding variant
        was a fresh bucket and the limit never fired."""
        client = APIClient()
        url = reverse("api-token-auth")

        with rates(login="2/minute"):
            for padded in ("victim", "victim ", " victim"):
                response = client.post(
                    url, {"username": padded, "password": "wrong"}, format="json"
                )

        self.assertEqual(response.status_code, status.HTTP_429_TOO_MANY_REQUESTS)

    def test_one_account_has_a_ceiling_across_every_address(self):
        """A guessing run spread over many addresses gets a fresh per-address
        budget each time; the per-account ceiling is what bounds it."""
        url = reverse("api-token-auth")

        with rates(login="2/minute", login_account="3/hour"):
            for index in range(3):
                APIClient(REMOTE_ADDR=f"198.51.100.{index}").post(
                    url, {"username": "victim", "password": "wrong"}, format="json"
                )
            response = APIClient(REMOTE_ADDR="198.51.100.99").post(
                url, {"username": "victim", "password": "wrong"}, format="json"
            )
            other = APIClient(REMOTE_ADDR="198.51.100.99").post(
                url, {"username": "someone_else", "password": "wrong"}, format="json"
            )

        self.assertEqual(response.status_code, status.HTTP_429_TOO_MANY_REQUESTS)
        self.assertNotEqual(other.status_code, status.HTTP_429_TOO_MANY_REQUESTS)

    def test_a_body_that_is_not_an_object_is_a_400_not_a_500(self):
        for body in (["victim", "wrong"], "victim"):
            with self.subTest(body=body):
                response = APIClient().post(
                    reverse("api-token-auth"), body, format="json"
                )
                self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_password_changes_are_throttled_per_user_not_per_address(self):
        """Behind one NAT every user shares an address. Keyed on the address
        and a username the body does not carry, one person's attempts used
        up everyone's - including anonymous sign-ins."""
        url = reverse("auth-password")
        other = make_user("throttle_neighbour", department="Registry")
        neighbour = APIClient()
        neighbour.force_authenticate(other)
        body = {"current_password": "wrong", "new_password": "N3w-pass-phrase!"}

        with rates(login="2/minute"):
            for _ in range(3):
                self.client.post(url, body, format="json")
            exhausted = self.client.post(url, body, format="json")
            response = neighbour.post(url, body, format="json")

        self.assertEqual(exhausted.status_code, status.HTTP_429_TOO_MANY_REQUESTS)
        self.assertNotEqual(response.status_code, status.HTTP_429_TOO_MANY_REQUESTS)

    def test_upload_is_throttled_but_reading_is_not(self):
        with rates(upload="1/hour"):
            first = self.client.post(
                reverse("document-list"),
                {"title": "x", "file": SimpleUploadedFile("a.txt", b"text")},
                format="multipart",
            )
            self.assertIn(first.status_code, (202, 400))

            second = self.client.post(
                reverse("document-list"),
                {"title": "y", "file": SimpleUploadedFile("b.txt", b"text")},
                format="multipart",
            )
            self.assertEqual(second.status_code, status.HTTP_429_TOO_MANY_REQUESTS)

            # Reading still works while the write budget is exhausted.
            self.assertEqual(
                self.client.get(reverse("document-list")).status_code,
                status.HTTP_200_OK,
            )


class SecretKeyGuardTests(SimpleTestCase):
    """Production settings refuse to load without a real secret key.

    Settings are read once per process, so each case loads them in a fresh
    interpreter, exactly as a deployment would.
    """

    def load_settings(self, key):
        import os
        import subprocess
        import sys
        from pathlib import Path

        return subprocess.run(
            [sys.executable, "-c", "import config.settings"],
            cwd=Path(__file__).resolve().parents[1],
            env={
                **os.environ,
                "DJANGO_DEBUG": "False",
                "DJANGO_SECRET_KEY": key,
                "DJANGO_ALLOWED_HOSTS": "docs.example.org",
            },
            capture_output=True,
            text=True,
            timeout=120,
        )

    def test_an_empty_key_is_refused(self):
        result = self.load_settings("")
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("ImproperlyConfigured", result.stderr)

    def test_the_shipped_placeholders_are_refused(self):
        """`.env.example` ships `change-me`. Deployed unedited, it passed the
        guard, and the deployment's key was public."""
        for key in ("change-me", "unsafe-development-key"):
            with self.subTest(key=key):
                result = self.load_settings(key)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("ImproperlyConfigured", result.stderr)

    def test_a_real_key_is_accepted(self):
        result = self.load_settings("k8Vq2-not-a-placeholder-" + "x" * 40)
        self.assertEqual(result.returncode, 0, result.stderr[-2000:])


class TransportSecurityTests(SimpleTestCase):
    """The production transport settings, asserted from inside the suite.

    `SECURE_SSL_REDIRECT` is switched off under the test runner, because
    Django's test client speaks plain HTTP and a 301 on every request makes
    the suite unrunnable (see config/settings.py for the full reasoning).
    That exemption is narrow, and it needs a guard: without one, the same
    line could later be edited to disable the redirect in production and no
    test in this repository would notice.

    So this reads the settings MODULE rather than the live configuration -
    the module is what a deployment loads, and it is the thing that would
    have been changed.
    """

    def source(self):
        from pathlib import Path

        import config.settings as settings_module

        return Path(settings_module.__file__).read_text(encoding="utf-8")

    def test_the_ssl_redirect_exemption_is_scoped_to_the_test_runner(self):
        source = self.source()

        self.assertIn("if RUNNING_TESTS:\n    SECURE_SSL_REDIRECT = False", source)
        self.assertIn('RUNNING_TESTS = "test" in sys.argv', source)

    def test_the_redirect_defaults_on_outside_debug(self):
        """`not DEBUG` is the default, so a production deployment that sets
        nothing still redirects."""
        self.assertIn(
            'SECURE_SSL_REDIRECT = os.getenv("SECURE_SSL_REDIRECT", str(not DEBUG))',
            self.source(),
        )

    def test_secure_cookies_are_not_exempted(self):
        """Only the redirect is lifted for tests. If someone widens the
        exemption to cookies, HSTS or the proxy header, the suite stops
        exercising the configuration a deployment actually uses.

        Compared against the MODULE's DEBUG, not the live setting. The test
        runner forces `settings.DEBUG = False` after the module has loaded,
        so on a developer machine with DJANGO_DEBUG=True the live value
        disagrees with the one the cookie flags were derived from, and the
        assertion failed on every laptop while passing in CI."""
        from django.conf import settings

        import config.settings as settings_module

        self.assertTrue(settings.RUNNING_TESTS, "these tests run under the runner")
        self.assertEqual(settings.SESSION_COOKIE_SECURE, not settings_module.DEBUG)
        self.assertEqual(settings.CSRF_COOKIE_SECURE, not settings_module.DEBUG)


class AdminMountTests(SimpleTestCase):
    """Django's admin must not own the application's /admin/ prefix.

    The edge (nginx in production, Vite's proxy in development) sends a path
    to Django or to the frontend by prefix. With Django's admin at /admin/,
    the application's administration screens - /admin/users, /admin/audit,
    /admin/system - loaded only through in-app navigation: a reload, a
    bookmark or a new tab was answered by Django's admin login, and a
    superuser got a 404.
    """

    def test_django_admin_lives_under_its_own_prefix(self):
        self.assertEqual(reverse("admin:index"), "/django-admin/")

    def test_the_application_admin_paths_are_not_django_routes(self):
        from django.urls import Resolver404, resolve

        for path in ("/admin/", "/admin/users", "/admin/audit", "/admin/system"):
            with self.subTest(path=path), self.assertRaises(Resolver404):
                resolve(path)
